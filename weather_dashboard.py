import requests
import pandas as pd
import streamlit as st
import plotly.express as px


# =========================
# PAGE CONFIGURATION
# =========================

st.set_page_config(
    page_title="Weather Analytics",
    page_icon="🌤️",
    layout="wide",
    initial_sidebar_state="expanded"
)


# =========================
# API CONFIGURATION
# =========================

GEOCODING_URL = "https://geocoding-api.open-meteo.com/v1/search"
WEATHER_URL = "https://api.open-meteo.com/v1/forecast"


# =========================
# CUSTOM DESIGN
# =========================

st.markdown("""
<style>
    .stApp {
        background-color: #0b1120;
        color: #f1f5f9;
    }

    [data-testid="stSidebar"] {
        background-color: #111827;
    }

    [data-testid="stMetric"] {
        background-color: #172033;
        border: 1px solid #263449;
        padding: 20px;
        border-radius: 16px;
    }

    [data-testid="stMetricLabel"] {
        color: #a5b4fc;
    }

    [data-testid="stMetricValue"] {
        color: #ffffff;
    }

    .main-title {
        font-size: 36px;
        font-weight: 800;
        color: #f8fafc;
        margin-bottom: 5px;
    }

    .subtitle {
        color: #94a3b8;
        font-size: 15px;
        margin-bottom: 25px;
    }

    .section-title {
        font-size: 22px;
        font-weight: 700;
        color: #e2e8f0;
        margin-top: 20px;
    }

    div.stButton > button {
        width: 100%;
        border-radius: 10px;
        height: 45px;
        font-weight: 700;
    }

    div[data-testid="stDataFrame"] {
        border: 1px solid #263449;
        border-radius: 12px;
    }
</style>
""", unsafe_allow_html=True)


# =========================
# GET CITY COORDINATES
# =========================

def get_city_coordinates(city_name):

    params = {
        "name": city_name,
        "count": 1,
        "language": "en",
        "format": "json"
    }

    response = requests.get(
        GEOCODING_URL,
        params=params,
        timeout=15
    )

    response.raise_for_status()

    data = response.json()
    results = data.get("results", [])

    if not results:
        raise ValueError(
            "City not found. Check the English spelling."
        )

    city = results[0]

    return {
        "name": city["name"],
        "country": city.get("country", "Unknown"),
        "latitude": city["latitude"],
        "longitude": city["longitude"]
    }


# =========================
# GET WEATHER DATA
# =========================

def get_weather_data(latitude, longitude):

    params = {
        "latitude": latitude,
        "longitude": longitude,
        "daily": [
            "weather_code",
            "temperature_2m_max",
            "temperature_2m_min",
            "precipitation_sum",
            "wind_speed_10m_max"
        ],
        "forecast_days": 7,
        "timezone": "auto"
    }

    response = requests.get(
        WEATHER_URL,
        params=params,
        timeout=20
    )

    response.raise_for_status()

    data = response.json()

    if "daily" not in data:
        raise ValueError("Weather data is unavailable.")

    return data["daily"]


# =========================
# CREATE DATAFRAME
# =========================

def create_weather_report(weather_data):

    df = pd.DataFrame({
        "Date": weather_data["time"],
        "Weather Code": weather_data["weather_code"],
        "Max Temperature": weather_data["temperature_2m_max"],
        "Min Temperature": weather_data["temperature_2m_min"],
        "Precipitation": weather_data["precipitation_sum"],
        "Max Wind Speed": weather_data["wind_speed_10m_max"]
    })

    df["Date"] = pd.to_datetime(df["Date"])

    return df


# =========================
# WEATHER CODE DESCRIPTION
# =========================

def get_weather_description(code):

    descriptions = {
        0: "Clear sky",
        1: "Mainly clear",
        2: "Partly cloudy",
        3: "Overcast",
        45: "Fog",
        48: "Depositing rime fog",
        51: "Light drizzle",
        53: "Moderate drizzle",
        55: "Dense drizzle",
        61: "Slight rain",
        63: "Moderate rain",
        65: "Heavy rain",
        71: "Slight snow",
        73: "Moderate snow",
        75: "Heavy snow",
        80: "Rain showers",
        81: "Moderate rain showers",
        82: "Violent rain showers",
        95: "Thunderstorm",
        96: "Thunderstorm with hail",
        99: "Thunderstorm with heavy hail"
    }

    return descriptions.get(int(code), "Unknown")


# =========================
# SIDEBAR
# =========================

with st.sidebar:

    st.markdown("## 🌤️ Weather Analytics")

    st.caption("Real-time API-powered dashboard")

    st.divider()

    city_name = st.text_input(
        "Search City",
        value="Tehran",
        placeholder="Enter city in English"
    )

    search_button = st.button(
        "🔍 Get Weather Forecast",
        type="primary"
    )

    st.divider()

    st.markdown("### About")

    st.caption(
        "This dashboard retrieves a 7-day forecast "
        "using the Open-Meteo weather API."
    )

    st.caption("Weather data: Open-Meteo")


# =========================
# MAIN HEADER
# =========================

st.markdown(
    '<div class="main-title">Weather Analytics Dashboard</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="subtitle">'
    'Explore weather forecasts, analyze temperature trends, '
    'and export reports.'
    '</div>',
    unsafe_allow_html=True
)


# =========================
# FETCH DATA
# =========================

if "weather_df" not in st.session_state:
    st.session_state.weather_df = None

if "city_info" not in st.session_state:
    st.session_state.city_info = None


if search_button:

    if not city_name.strip():

        st.warning("Please enter a city name.")

    else:

        try:

            with st.spinner("Fetching weather data..."):

                city = get_city_coordinates(city_name.strip())

                weather_data = get_weather_data(
                    city["latitude"],
                    city["longitude"]
                )

                df = create_weather_report(weather_data)

                st.session_state.weather_df = df
                st.session_state.city_info = city

        except requests.exceptions.Timeout:

            st.error("The API request timed out. Try again.")

        except requests.exceptions.ConnectionError:

            st.error(
                "Could not connect to the API. "
                "Check your internet connection."
            )

        except requests.exceptions.HTTPError as error:

            st.error(f"HTTP error: {error}")

        except requests.exceptions.RequestException as error:

            st.error(f"API request failed: {error}")

        except (ValueError, KeyError, TypeError) as error:

            st.error(f"Data error: {error}")


# =========================
# DISPLAY DASHBOARD
# =========================

df = st.session_state.weather_df
city = st.session_state.city_info


if df is not None and city is not None:

    st.markdown(
        f"## 📍 {city['name']}, {city['country']}"
    )

    st.caption(
        f"Coordinates: {city['latitude']:.3f}, "
        f"{city['longitude']:.3f}"
    )

    # ---------------------
    # KPI CARDS
    # ---------------------

    avg_max = df["Max Temperature"].mean()
    avg_min = df["Min Temperature"].mean()
    highest = df["Max Temperature"].max()
    total_rain = df["Precipitation"].sum()

    col1, col2, col3, col4 = st.columns(4)

    col1.metric(
        "🌡️ Avg. High",
        f"{avg_max:.1f} °C"
    )

    col2.metric(
        "❄️ Avg. Low",
        f"{avg_min:.1f} °C"
    )

    col3.metric(
        "🔥 Highest Temperature",
        f"{highest:.1f} °C"
    )

    col4.metric(
        "🌧️ Total Precipitation",
        f"{total_rain:.1f} mm"
    )

    st.divider()

    # ---------------------
    # TEMPERATURE CHART
    # ---------------------

    st.markdown(
        '<div class="section-title">Temperature Trends</div>',
        unsafe_allow_html=True
    )

    temperature_data = df.melt(
        id_vars=["Date"],value_vars=[
            "Max Temperature",
            "Min Temperature"
        ],
        var_name="Temperature Type",
        value_name="Temperature (°C)"
    )

    fig_temp = px.line(
        temperature_data,
        x="Date",
        y="Temperature (°C)",
        color="Temperature Type",
        markers=True,
        template="plotly_dark",
        labels={
            "Date": "Date",
            "Temperature (°C)": "Temperature (°C)"
        }
    )

    fig_temp.update_layout(
        paper_bgcolor="#0b1120",
        plot_bgcolor="#111827",
        legend_title_text="",
        margin=dict(l=10, r=10, t=25, b=10),
        hovermode="x unified"
    )

    st.plotly_chart(
        fig_temp,
        use_container_width=True
    )

    # ---------------------
    # PRECIPITATION AND WIND
    # ---------------------

    col_left, col_right = st.columns(2)

    with col_left:

        st.markdown("### 🌧️ Daily Precipitation")

        fig_rain = px.bar(
            df,
            x="Date",
            y="Precipitation",
            template="plotly_dark",
            labels={
                "Date": "Date",
                "Precipitation": "Precipitation (mm)"
            }
        )

        fig_rain.update_layout(
            paper_bgcolor="#0b1120",
            plot_bgcolor="#111827",
            margin=dict(l=10, r=10, t=20, b=10)
        )

        st.plotly_chart(
            fig_rain,
            use_container_width=True
        )

    with col_right:

        st.markdown("### 💨 Maximum Wind Speed")

        fig_wind = px.bar(
            df,
            x="Date",
            y="Max Wind Speed",
            template="plotly_dark",
            labels={
                "Date": "Date",
                "Max Wind Speed": "Wind Speed (km/h)"
            }
        )

        fig_wind.update_layout(
            paper_bgcolor="#0b1120",
            plot_bgcolor="#111827",
            margin=dict(l=10, r=10, t=20, b=10)
        )

        st.plotly_chart(
            fig_wind,
            use_container_width=True
        )

    # ---------------------
    # WEATHER TABLE
    # ---------------------

    st.markdown(
        '<div class="section-title">7-Day Forecast</div>',
        unsafe_allow_html=True
    )

    display_df = df.copy()

    display_df["Weather"] = display_df["Weather Code"].apply(
        get_weather_description
    )

    display_df = display_df[
        [
            "Date",
            "Weather",
            "Max Temperature",
            "Min Temperature",
            "Precipitation",
            "Max Wind Speed"
        ]
    ]

    display_df["Date"] = display_df["Date"].dt.strftime("%Y-%m-%d")

    display_df = display_df.rename(columns={
        "Date": "Date",
        "Weather": "Conditions",
        "Max Temperature": "Max Temp (°C)",
        "Min Temperature": "Min Temp (°C)",
        "Precipitation": "Precipitation (mm)",
        "Max Wind Speed": "Wind Speed (km/h)"
    })

    st.dataframe(
        display_df,
        use_container_width=True,
        hide_index=True
    )

# =========================
    # PROFESSIONAL EXCEL REPORT
    # =========================

    from io import BytesIO
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    # ---------------------
    # SUMMARY DATA
    # ---------------------

    warmest_day = df.loc[df["Max Temperature"].idxmax()]
    coldest_day = df.loc[df["Min Temperature"].idxmin()]

    summary_df = pd.DataFrame({
        "Metric": [
            "City",
            "Country",
            "Forecast Days",
            "Average Maximum Temperature (°C)",
            "Average Minimum Temperature (°C)",
            "Highest Temperature (°C)",
            "Lowest Temperature (°C)",
            "Total Precipitation (mm)",
            "Average Wind Speed (km/h)",
            "Maximum Wind Speed (km/h)",
            "Warmest Day",
            "Coldest Day"
        ],
        "Value": [
            city["name"],
            city["country"],
            len(df),
            round(df["Max Temperature"].mean(), 1),
            round(df["Min Temperature"].mean(), 1),
            round(df["Max Temperature"].max(), 1),
            round(df["Min Temperature"].min(), 1),
            round(df["Precipitation"].sum(), 1),
            round(df["Max Wind Speed"].mean(), 1),
            round(df["Max Wind Speed"].max(), 1),
            warmest_day["Date"].strftime("%Y-%m-%d"),
            coldest_day["Date"].strftime("%Y-%m-%d")
        ]
    })

    # ---------------------
    # CREATE EXCEL FILE
    # ---------------------

    excel_buffer = BytesIO()

    with pd.ExcelWriter(
        excel_buffer,
        engine="openpyxl"
    ) as writer:

        # Sheet 1: Daily weather data
        display_df.to_excel(
            writer,
            sheet_name="Weather Data",
            index=False
        )

        # Sheet 2: Statistical summary
        summary_df.to_excel(
            writer,
            sheet_name="Summary",
            index=False
        )

        # ---------------------
        # FORMAT WORKSHEETS
        # ---------------------

        for worksheet in writer.sheets.values():

            # Header formatting
            for cell in worksheet[1]:

                cell.font = Font(
                    bold=True,
                    color="FFFFFF",
                    size=11
                )

                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="17365D"
                )

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            worksheet.sheet_view.showGridLines = False
            worksheet.row_dimensions[1].height = 25

            # Automatic column widths
            for column_cells in worksheet.columns:

                max_length = max(
                    len(str(cell.value or ""))
                    for cell in column_cells
                )

                column_letter = get_column_letter(
                    column_cells[0].column
                )

                worksheet.column_dimensions[
                    column_letter
                ].width = min(max_length + 4, 40)

    # Get the completed Excel file
    excel_data = excel_buffer.getvalue()

    # ---------------------
    # DOWNLOAD EXCEL
    # ---------------------

    st.download_button(
        label="📥 Download Professional Excel Report",
        data=excel_data,
        file_name="weather_report.xlsx",
        mime=(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        ),
        type="primary"
    )

    # ---------------------
    # SAVE EXCEL LOCALLY
    # ---------------------

    with open("weather_report.xlsx", "wb") as file:
        file.write(excel_data)

    st.caption(
        "Excel report contains two worksheets: "
        "Weather Data and Summary."
    )